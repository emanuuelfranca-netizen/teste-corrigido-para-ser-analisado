import os, io
from datetime import datetime, date, timedelta
from pathlib import Path
from flask import Flask, render_template, request, redirect, url_for, flash, send_file, jsonify, abort
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from sqlalchemy import text
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.graphics.shapes import Drawing, String
from reportlab.graphics.charts.piecharts import Pie
from openpyxl import Workbook
from docx import Document
BASE=Path(__file__).parent
app=Flask(__name__); app.config['SECRET_KEY']=os.getenv('SECRET_KEY','dev-change-me'); app.config['SQLALCHEMY_DATABASE_URI']=os.getenv('DATABASE_URL','sqlite:///'+str(BASE/'fivecom.db')); app.config['SQLALCHEMY_TRACK_MODIFICATIONS']=False; app.config['UPLOAD_FOLDER']=str(BASE/'static/uploads'); Path(app.config['UPLOAD_FOLDER']).mkdir(parents=True,exist_ok=True)
db=SQLAlchemy(app); login=LoginManager(app); login.login_view='login_page'; login.login_message='Faça login para continuar.'
product_suppliers=db.Table('product_suppliers',db.Column('product_id',db.Integer,db.ForeignKey('product.id'),primary_key=True),db.Column('supplier_id',db.Integer,db.ForeignKey('supplier.id'),primary_key=True))
class User(UserMixin,db.Model):
 id=db.Column(db.Integer,primary_key=True); name=db.Column(db.String(120),nullable=False); username=db.Column(db.String(80),unique=True,nullable=False); password=db.Column(db.String(255),nullable=False); role=db.Column(db.String(30),default='Operador'); active=db.Column(db.Boolean,default=True); permissions=db.Column(db.Text,default='')
 def can(self,m): return self.role=='Administrador' or m in (self.permissions or '').split(',')
class Setting(db.Model):
 id=db.Column(db.Integer,primary_key=True); system_name=db.Column(db.String(120),default='FiveCom ERP'); company_name=db.Column(db.String(120),default='Estoque Enterprise'); browser_title=db.Column(db.String(120),default='FiveCom ERP'); primary_color=db.Column(db.String(20),default='#2563eb'); secondary_color=db.Column(db.String(20),default='#0f172a'); logo=db.Column(db.String(255)); login_logo=db.Column(db.String(255)); favicon=db.Column(db.String(255)); login_position=db.Column(db.String(20),default='center'); footer=db.Column(db.String(255),default='Documento gerado automaticamente pelo FiveCom ERP'); disabled_modules=db.Column(db.Text,default='')
class Category(db.Model):
 id=db.Column(db.Integer,primary_key=True); name=db.Column(db.String(100),unique=True,nullable=False); description=db.Column(db.Text); active=db.Column(db.Boolean,default=True)
class Supplier(db.Model):
 id=db.Column(db.Integer,primary_key=True); name=db.Column(db.String(150),nullable=False); document=db.Column(db.String(40)); address=db.Column(db.String(255)); email=db.Column(db.String(120)); phone=db.Column(db.String(40)); notes=db.Column(db.Text); created_at=db.Column(db.DateTime,default=datetime.now); products=db.relationship('Product',secondary=product_suppliers,back_populates='suppliers')
class Product(db.Model):
 id=db.Column(db.Integer,primary_key=True); code=db.Column(db.String(50),unique=True,nullable=False); name=db.Column(db.String(150),nullable=False); category_id=db.Column(db.Integer,db.ForeignKey('category.id')); category=db.relationship('Category'); supplier_id=db.Column(db.Integer,db.ForeignKey('supplier.id')); unit=db.Column(db.String(20),default='UN'); cost=db.Column(db.Float,default=0); price=db.Column(db.Float,default=0); stock=db.Column(db.Float,default=0); min_stock=db.Column(db.Float,default=0); max_stock=db.Column(db.Float,default=0); location=db.Column(db.String(100)); perishable=db.Column(db.Boolean,default=False); batch=db.Column(db.String(80)); manufacture_date=db.Column(db.Date); expiry_date=db.Column(db.Date); attachment=db.Column(db.String(255)); photo=db.Column(db.String(255)); active=db.Column(db.Boolean,default=True); suppliers=db.relationship('Supplier',secondary=product_suppliers,back_populates='products')
 @property
 def days_to_expiry(self): return (self.expiry_date-date.today()).days if self.expiry_date else None
class Movement(db.Model):
 id=db.Column(db.Integer,primary_key=True); product_id=db.Column(db.Integer,db.ForeignKey('product.id')); kind=db.Column(db.String(20)); quantity=db.Column(db.Float); reason=db.Column(db.String(100)); created_at=db.Column(db.DateTime,default=datetime.now); user_id=db.Column(db.Integer); product=db.relationship('Product')
class Task(db.Model):
 id=db.Column(db.Integer,primary_key=True); title=db.Column(db.String(200)); description=db.Column(db.Text); due_date=db.Column(db.Date); priority=db.Column(db.String(20),default='Média'); status=db.Column(db.String(30),default='Pendente'); user_id=db.Column(db.Integer,db.ForeignKey('user.id')); user=db.relationship('User')
class Asset(db.Model):
 id=db.Column(db.Integer,primary_key=True); name=db.Column(db.String(160)); tag=db.Column(db.String(80),unique=True); serial=db.Column(db.String(100)); location=db.Column(db.String(100)); responsible=db.Column(db.String(120)); value=db.Column(db.Float,default=0); status=db.Column(db.String(30),default='Ativo')
class Knowledge(db.Model):
 id=db.Column(db.Integer,primary_key=True); title=db.Column(db.String(200)); kind=db.Column(db.String(40)); content=db.Column(db.Text); supplier_id=db.Column(db.Integer,db.ForeignKey('supplier.id')); supplier=db.relationship('Supplier'); created_at=db.Column(db.DateTime,default=datetime.now)
class Automation(db.Model):
 id=db.Column(db.Integer,primary_key=True); name=db.Column(db.String(150)); entity=db.Column(db.String(80)); trigger=db.Column(db.String(120)); action=db.Column(db.String(120)); active=db.Column(db.Boolean,default=True)
class ChartConfig(db.Model):
 id=db.Column(db.Integer,primary_key=True); name=db.Column(db.String(150)); entity=db.Column(db.String(80)); metric=db.Column(db.String(80)); chart_type=db.Column(db.String(30)); period=db.Column(db.String(30))
class Audit(db.Model):
 id=db.Column(db.Integer,primary_key=True); user=db.Column(db.String(120)); action=db.Column(db.String(255)); created_at=db.Column(db.DateTime,default=datetime.now); ip=db.Column(db.String(80))
@login.user_loader
def load_user(i): return db.session.get(User,int(i))
def settings():
 s=Setting.query.first()
 if not s:s=Setting();db.session.add(s);db.session.commit()
 return s
@app.context_processor
def ctx():
 s=settings(); return dict(cfg=s,module_enabled=lambda m:m not in (s.disabled_modules or '').split(','),now=datetime.now())
def audit(a): db.session.add(Audit(user=current_user.name if current_user.is_authenticated else 'Sistema',action=a,ip=request.headers.get('X-Forwarded-For',request.remote_addr)));db.session.commit()
def upload(n):
 f=request.files.get(n)
 if not f or not f.filename:return None
 x=datetime.now().strftime('%Y%m%d%H%M%S%f_')+secure_filename(f.filename);f.save(Path(app.config['UPLOAD_FOLDER'])/x);return x
def category_totals():
 out={}
 for p in Product.query.filter_by(active=True).all(): out[p.category.name if p.category else 'Sem categoria']=out.get(p.category.name if p.category else 'Sem categoria',0)+p.stock
 return out
def product_totals(): return {p.name:p.stock for p in Product.query.filter_by(active=True).all() if p.stock}
def pie_drawing(data,title):
 d=Drawing(500,260); d.add(String(10,240,title,fontSize=13)); pie=Pie();pie.x=120;pie.y=25;pie.width=180;pie.height=180;pie.data=list(data.values()) or [1];pie.labels=[f'{k}: {v:g}' for k,v in data.items()] or ['Sem dados']; palette=[colors.HexColor(x) for x in ['#2563eb','#16a34a','#f59e0b','#dc2626','#7c3aed','#0891b2','#db2777','#65a30d']]
 for i in range(len(pie.data)):pie.slices[i].fillColor=palette[i%len(palette)]
 d.add(pie);return d
@app.route('/login',methods=['GET','POST'])
def login_page():
 if request.method=='POST':
  u=User.query.filter_by(username=request.form['username'],active=True).first()
  if u and check_password_hash(u.password,request.form['password']):login_user(u);audit('Login realizado');return redirect(url_for('home'))
  flash('Usuário ou senha inválidos.','error')
 return render_template('login.html')
@app.route('/logout')
@login_required
def logout():logout_user();return redirect(url_for('login_page'))
@app.route('/')
@login_required
def home():
 today=date.today();return render_template('home.html',tasks=Task.query.filter_by(user_id=current_user.id).filter(Task.due_date<=today).order_by(Task.status).all(),low=Product.query.filter(Product.stock<=0,Product.active==True).count(),expiring=Product.query.filter(Product.expiry_date!=None,Product.expiry_date<=today+timedelta(days=30),Product.active==True).all(),movements=Movement.query.order_by(Movement.created_at.desc()).limit(8).all(),products=Product.query.filter_by(active=True).count(),value=sum(p.stock*p.cost for p in Product.query.filter_by(active=True)))
@app.route('/products',methods=['GET','POST'])
@login_required
def products():
 if request.method=='POST':
  per=request.form.get('perishable')=='on';p=Product(code=request.form['code'].strip(),name=request.form['name'],category_id=request.form.get('category_id') or None,unit=request.form.get('unit','UN'),cost=float(request.form.get('cost') or 0),price=float(request.form.get('price') or 0),stock=float(request.form.get('stock') or 0),location=request.form.get('location'),perishable=per,batch=request.form.get('batch') if per else None,manufacture_date=date.fromisoformat(request.form['manufacture_date']) if per and request.form.get('manufacture_date') else None,expiry_date=date.fromisoformat(request.form['expiry_date']) if per and request.form.get('expiry_date') else None,attachment=upload('attachment'),photo=upload('photo'))
  p.suppliers=Supplier.query.filter(Supplier.id.in_([int(x) for x in request.form.getlist('supplier_ids')])).all() if request.form.getlist('supplier_ids') else []
  db.session.add(p);db.session.commit();audit('Produto criado: '+p.name);flash('Produto salvo com sucesso.','success');return redirect(url_for('products'))
 q=Product.query;cat=request.args.get('category');status=request.args.get('status','active')
 if cat:q=q.filter_by(category_id=int(cat))
 q=q.filter_by(active=(status!='inactive'))
 return render_template('products.html',items=q.order_by(Product.name).all(),suppliers=Supplier.query.order_by(Supplier.name).all(),categories=Category.query.filter_by(active=True).order_by(Category.name).all(),selected_cat=cat,status=status,open_form=request.args.get('new')=='1')
@app.route('/products/<int:i>/status',methods=['POST'])
@login_required
def product_status(i):
 p=db.session.get(Product,i) or abort(404);p.active=not p.active;db.session.commit();audit(('Reativado: ' if p.active else 'Desativado: ')+p.name);flash('Produto reativado.' if p.active else 'Produto desativado. Ele continua disponível no filtro de inativos.','success');return redirect(url_for('products',status='active' if p.active else 'inactive'))
@app.route('/categories',methods=['GET','POST'])
@login_required
def categories():
 if request.method=='POST':
  db.session.add(Category(name=request.form['name'],description=request.form.get('description')));db.session.commit();audit('Categoria criada: '+request.form['name']);return redirect(url_for('categories'))
 return render_template('categories.html',items=Category.query.order_by(Category.name).all())
@app.route('/movements',methods=['GET','POST'])
@login_required
def movements():
 if request.method=='POST':
  p=db.session.get(Product,int(request.form['product_id']));q=float(request.form['quantity']);kind=request.form['kind']
  if kind=='Saída' and q>p.stock:flash('Saldo insuficiente.','error');return redirect(url_for('movements'))
  p.stock+=q if kind=='Entrada' else -q;db.session.add(Movement(product=p,kind=kind,quantity=q,reason=request.form.get('reason'),user_id=current_user.id));db.session.commit();audit(f'{kind} de {q} em {p.name}');return redirect(url_for('movements'))
 return render_template('movements.html',items=Movement.query.order_by(Movement.created_at.desc()).all(),products=Product.query.filter_by(active=True).all())
@app.route('/tasks',methods=['GET','POST'])
@login_required
def tasks():
 if request.method=='POST':db.session.add(Task(title=request.form['title'],description=request.form.get('description'),due_date=date.fromisoformat(request.form['due_date']),priority=request.form['priority'],user_id=int(request.form.get('user_id') or current_user.id)));db.session.commit();audit('Tarefa criada');return redirect(url_for('tasks'))
 return render_template('tasks.html',items=Task.query.order_by(Task.due_date).all(),users=User.query.filter_by(active=True).all())
@app.route('/tasks/<int:i>/toggle',methods=['POST'])
@login_required
def task_toggle(i):
 t=db.session.get(Task,i);t.status='Concluída' if t.status!='Concluída' else 'Pendente';db.session.commit();return redirect(request.referrer or url_for('tasks'))
@app.route('/suppliers',methods=['GET','POST'])
@login_required
def suppliers():
 if request.method=='POST':
  s=Supplier(name=request.form['name'],document=request.form.get('document'),address=request.form.get('address'),email=request.form.get('email'),phone=request.form.get('phone'),notes=request.form.get('notes'));s.products=Product.query.filter(Product.id.in_([int(x) for x in request.form.getlist('product_ids')])).all() if request.form.getlist('product_ids') else [];db.session.add(s);db.session.commit();audit('Fornecedor criado: '+s.name);return redirect(url_for('supplier_detail',i=s.id))
 return render_template('suppliers.html',items=Supplier.query.order_by(Supplier.name).all(),products=Product.query.filter_by(active=True).all())
@app.route('/suppliers/<int:i>',methods=['GET','POST'])
@login_required
def supplier_detail(i):
 s=db.session.get(Supplier,i) or abort(404)
 if request.method=='POST':db.session.add(Knowledge(title=request.form['title'],kind='Relacionamento',content=request.form['content'],supplier=s));db.session.commit();audit('Nota adicionada ao fornecedor: '+s.name);return redirect(url_for('supplier_detail',i=i))
 return render_template('supplier_detail.html',s=s,notes=Knowledge.query.filter_by(supplier_id=i).order_by(Knowledge.created_at.desc()).all())
@app.route('/assets',methods=['GET','POST'])
@login_required
def assets():
 if request.method=='POST':db.session.add(Asset(name=request.form['name'],tag=request.form['tag'],serial=request.form.get('serial'),location=request.form.get('location'),responsible=request.form.get('responsible'),value=float(request.form.get('value') or 0),status=request.form.get('status')));db.session.commit();return redirect(url_for('assets'))
 return render_template('assets.html',items=Asset.query.all())
@app.route('/knowledge',methods=['GET','POST'])
@login_required
def knowledge():
 if request.method=='POST':db.session.add(Knowledge(title=request.form['title'],kind=request.form['kind'],content=request.form['content'],supplier_id=request.form.get('supplier_id') or None));db.session.commit();return redirect(url_for('knowledge'))
 return render_template('knowledge.html',items=Knowledge.query.all(),suppliers=Supplier.query.all())
@app.route('/automations',methods=['GET','POST'])
@login_required
def automations():
 if request.method=='POST':db.session.add(Automation(name=request.form['name'],entity=request.form['entity'],trigger=request.form['trigger'],action=request.form['action']));db.session.commit();return redirect(url_for('automations'))
 return render_template('automations.html',items=Automation.query.all())
@app.route('/charts',methods=['GET','POST'])
@login_required
def charts():
 if request.method=='POST':db.session.add(ChartConfig(name=request.form['name'],entity=request.form['entity'],metric=request.form['metric'],chart_type=request.form['chart_type'],period=request.form['period']));db.session.commit();return redirect(url_for('charts'))
 return render_template('charts.html',items=ChartConfig.query.all(),category_data=category_totals(),product_data=product_totals())
@app.route('/settings',methods=['GET','POST'])
@login_required
def config():
 if current_user.role!='Administrador':return redirect(url_for('home'))
 s=settings()
 if request.method=='POST':
  for x in ['system_name','company_name','browser_title','primary_color','secondary_color','login_position','footer']:setattr(s,x,request.form.get(x))
  s.disabled_modules=','.join(request.form.getlist('disabled_modules'))
  for f in ['logo','login_logo','favicon']:
   u=upload(f)
   if u:setattr(s,f,u)
  db.session.commit();audit('Configurações atualizadas');return redirect(url_for('config'))
 return render_template('settings.html',s=s)
@app.route('/users',methods=['GET','POST'])
@login_required
def users():
 if current_user.role!='Administrador':return redirect(url_for('home'))
 if request.method=='POST':db.session.add(User(name=request.form['name'],username=request.form['username'],password=generate_password_hash(request.form['password']),role=request.form['role'],permissions=','.join(request.form.getlist('permissions'))));db.session.commit();return redirect(url_for('users'))
 return render_template('users.html',items=User.query.all())
@app.route('/audits')
@login_required
def audits():return render_template('audits.html',items=Audit.query.order_by(Audit.created_at.desc()).limit(300).all())
@app.route('/reports',methods=['GET','POST'])
@login_required
def reports():
 if request.method=='POST':
  fmt=request.form['format'];rows=[[m.created_at.strftime('%d/%m/%Y'),m.product.name,m.kind,str(m.quantity),m.reason or ''] for m in Movement.query.order_by(Movement.created_at.desc()).all() if (m.kind=='Entrada' and request.form.get('entries')) or (m.kind=='Saída' and request.form.get('outputs'))];s=settings();title='Relatório de balanço de estoque';cats=category_totals();prods=product_totals();want=request.form.get('chart');display=request.form.get('display','ambos')
  if fmt=='pdf':
   b=io.BytesIO();doc=SimpleDocTemplate(b,pagesize=A4,rightMargin=36,leftMargin=36);st=getSampleStyleSheet();story=[]
   if s.logo and (Path(app.config['UPLOAD_FOLDER'])/s.logo).exists():story.append(Image(str(Path(app.config['UPLOAD_FOLDER'])/s.logo),width=120,height=50))
   story += [Paragraph(title.upper(),st['Title']),Spacer(1,8),Paragraph('Data de emissão: '+datetime.now().strftime('%d/%m/%Y %H:%M'),st['Normal']),Paragraph('Emitido por: '+current_user.name,st['Normal']),Spacer(1,15)]
   if want and display!='lista':story += [pie_drawing(cats,'Estoque por categoria'),pie_drawing(prods,'Estoque por produto')]
   if display!='grafico':
    data=[['Data','Produto','Tipo','Quantidade','Motivo']]+rows;tab=Table(data,repeatRows=1,colWidths=[65,150,60,70,120]);tab.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor(s.primary_color)),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.3,colors.grey),('FONTSIZE',(0,0),(-1,-1),8)]));story.append(tab)
   story += [Spacer(1,15),Paragraph(s.footer or '',st['Normal'])];doc.build(story);b.seek(0);return send_file(b,mimetype='application/pdf',as_attachment=True,download_name='relatorio_estoque.pdf')
  if fmt=='xlsx':
   wb=Workbook();ws=wb.active;ws.title='Balanço';ws.append([title]);ws.append(['Categoria','Saldo']);[ws.append([k,v]) for k,v in cats.items()];ws.append([]);ws.append(['Data','Produto','Tipo','Quantidade','Motivo']);[ws.append(r) for r in rows];b=io.BytesIO();wb.save(b);b.seek(0);return send_file(b,as_attachment=True,download_name='relatorio_estoque.xlsx')
  d=Document();d.add_heading(title,0);d.add_paragraph('Data de emissão: '+datetime.now().strftime('%d/%m/%Y %H:%M'));d.add_heading('Estoque por categoria',1);[d.add_paragraph(f'{k}: {v:g} unidades') for k,v in cats.items()];t=d.add_table(rows=1,cols=5);[setattr(c,'text',v) for c,v in zip(t.rows[0].cells,['Data','Produto','Tipo','Quantidade','Motivo'])]
  for r in rows:
   c=t.add_row().cells
   for a,v in zip(c,r):a.text=str(v)
  b=io.BytesIO();d.save(b);b.seek(0);return send_file(b,as_attachment=True,download_name='relatorio_estoque.docx')
 return render_template('reports.html')
@app.route('/api/chart/<kind>')
@login_required
def chart_api(kind):
 d=category_totals() if kind=='category' else product_totals();return jsonify(labels=list(d),values=list(d.values()))

def migrate():
 db.create_all();
 # Compatibilidade com bancos anteriores SQLite/PostgreSQL: adiciona colunas novas quando necessário.
 for table,col,typ in [('supplier','address','VARCHAR(255)'),('supplier','notes','TEXT'),('product','category_id','INTEGER'),('product','photo','VARCHAR(255)'),('knowledge','supplier_id','INTEGER')]:
  try:db.session.execute(text(f'ALTER TABLE {table} ADD COLUMN {col} {typ}'));db.session.commit()
  except Exception:db.session.rollback()
def seed():
 if User.query.count()==0:db.session.add(User(name='Administrador',username='admin',password=generate_password_hash('Admin@123'),role='Administrador',permissions='produtos,movimentacoes,tarefas,fornecedores,patrimonio,conhecimento,automacoes,graficos,relatorios'))
 if Setting.query.count()==0:db.session.add(Setting())
 db.session.commit()
 if Category.query.count()==0:
  cats=[Category(name='Alimentos'),Category(name='Bebidas'),Category(name='Higiene'),Category(name='Limpeza'),Category(name='Eletrônicos')];db.session.add_all(cats);db.session.commit()
 if Supplier.query.count()==0:
  db.session.add_all([Supplier(name='Distribuidora Nordeste',document='12.345.678/0001-90',address='Av. Principal, 500, Recife/PE',email='vendas@dnordeste.example',phone='(81) 3333-1000',notes='Entrega semanal'),Supplier(name='Fábrica Boa Compra',document='98.765.432/0001-10',address='Rua Industrial, 80, Jaboatão/PE',email='comercial@boacompra.example',phone='(81) 3222-2000')]);db.session.commit()
 if Product.query.count()==0:
  c={x.name:x for x in Category.query.all()};s=Supplier.query.all();ps=[Product(code='789100000001',name='Arroz 5 kg',category=c['Alimentos'],stock=48,cost=21,price=29,unit='UN',location='A1'),Product(code='789100000002',name='Feijão 1 kg',category=c['Alimentos'],stock=35,cost=7,price=10,unit='UN',location='A2'),Product(code='789100000003',name='Leite integral',category=c['Bebidas'],stock=24,cost=4.5,price=6.5,perishable=True,batch='LT-1026',manufacture_date=date.today()-timedelta(days=7),expiry_date=date.today()+timedelta(days=20)),Product(code='789100000004',name='Suco de uva',category=c['Bebidas'],stock=19,cost=8,price=12),Product(code='789100000005',name='Sabonete',category=c['Higiene'],stock=60,cost=2.2,price=4),Product(code='789100000006',name='Detergente',category=c['Limpeza'],stock=42,cost=2,price=3.5),Product(code='789100000007',name='Mouse USB',category=c['Eletrônicos'],stock=8,cost=25,price=45)];db.session.add_all(ps);db.session.flush();s[0].products=ps[:6];s[1].products=[ps[6]];db.session.commit()
 if Movement.query.count()==0:
  ps=Product.query.all();
  for n,p in enumerate(ps):db.session.add(Movement(product=p,kind='Entrada',quantity=10+n*2,reason='Compra',created_at=datetime.now()-timedelta(days=n),user_id=1))
  for p in ps[:4]:db.session.add(Movement(product=p,kind='Saída',quantity=3,reason='Venda',created_at=datetime.now()-timedelta(hours=p.id),user_id=1))
 if Task.query.count()==0:db.session.add_all([Task(title='Conferir lote de leite',due_date=date.today(),priority='Alta',user_id=1),Task(title='Revisar saldo de eletrônicos',due_date=date.today(),priority='Média',user_id=1)])
 if Asset.query.count()==0:db.session.add_all([Asset(name='Notebook Dell',tag='PAT-001',serial='SN1001',location='Administrativo',responsible='Administrador',value=3500),Asset(name='Leitor de código de barras',tag='PAT-002',serial='SN1002',location='Estoque',responsible='Operador',value=450)])
 if Knowledge.query.count()==0:db.session.add_all([Knowledge(title='Política de recebimento',kind='Procedimento',content='Conferir quantidade, lote e integridade antes de confirmar a entrada.'),Knowledge(title='Relacionamento Distribuidora Nordeste',kind='Relacionamento',content='Contato prefere pedidos às terças-feiras.',supplier_id=Supplier.query.first().id)])
 if Automation.query.count()==0:db.session.add(Automation(name='Aviso de validade',entity='Produtos',trigger='Validade próxima',action='Criar notificação'))
 if ChartConfig.query.count()==0:db.session.add_all([ChartConfig(name='Estoque por categoria',entity='Produtos',metric='Quantidade em estoque',chart_type='Pizza',period='Atual'),ChartConfig(name='Estoque por produto',entity='Produtos',metric='Quantidade em estoque',chart_type='Barras',period='Atual')])
 db.session.commit()
with app.app_context():migrate();seed()
if __name__=='__main__':app.run(host='0.0.0.0',port=int(os.getenv('PORT',8000)),debug=os.getenv('FLASK_DEBUG')=='1')
